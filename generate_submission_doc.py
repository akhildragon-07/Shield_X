from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def create_submission_document(output_path="Email_Deliverability_Checker_Submission.docx"):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    primary_color = RGBColor(79, 70, 229)
    dark_color = RGBColor(15, 23, 42)
    gray_color = RGBColor(100, 116, 139)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Email Deliverability Checker")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(16)
    run_sub = subtitle_p.add_run("Software Development Course Project - Assignment Submission Document")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = gray_color

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(2.2), Inches(4.3)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    meta_items = [
        ("Project Title:", "Email Deliverability Checker"),
        ("Course / Assignment:", "Software Development Course Project"),
        ("Live Application URL:", "[Pending Deployment - e.g., https://email-deliverability-checker.onrender.com]"),
        ("GitHub Repository URL:", "[Pending Publication - e.g., https://github.com/<username>/email-deliverability-checker]"),
        ("Tech Stack:", "Python 3.12, Flask, dnspython, Gunicorn, Vanilla JS, CSS3, HTML5")
    ]

    for row_idx, (label, val) in enumerate(meta_items):
        cell_lbl = table.rows[row_idx].cells[0]
        cell_val = table.rows[row_idx].cells[1]

        set_cell_background(cell_lbl, "F8FAFC")
        set_cell_background(cell_val, "FFFFFF")

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(4)
        p_lbl.paragraph_format.space_after = Pt(4)
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.color.rgb = dark_color

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(4)
        p_val.paragraph_format.space_after = Pt(4)
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)
        if "[" in val:
            r_val.font.italic = True
            r_val.font.color.rgb = RGBColor(194, 65, 12)
        else:
            r_val.font.color.rgb = dark_color

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Project Summary")
    r_h1.font.color.rgb = primary_color
    r_h1.font.bold = True

    p_sum = doc.add_paragraph(
        "I built this Email Deliverability Checker to help users check whether an email list is safe to send "
        "to without actually sending test emails or using paid third-party APIs. Users can paste up to 100 emails at a time. "
        "The application cleans duplicate entries, checks the email syntax, resolves domain MX records through DNS, and "
        "attempts a safe SMTP handshake.\n\n"
        "To avoid misleading users, the app categorizes results into three realistic statuses: 'Likely deliverable', "
        "'Likely to bounce', and 'Unknown'. All processing happens in memory without storing email addresses in a database."
    )
    p_sum.paragraph_format.space_after = Pt(12)

    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Main Features")
    r_h2.font.color.rgb = primary_color
    r_h2.font.bold = True

    features = [
        ("Multi-step checking logic: ", "Combines regex format validation, dnspython for MX and fallback A record lookups, and optional socket/SMTP recipient probes."),
        ("No actual emails sent: ", "Only performs the recipient check up to the RCPT TO command and closes the connection with QUIT. Never sends DATA or message bodies."),
        ("Batch processing & deduplication: ", "Automatically strips whitespace and duplicate addresses, with a safety cap of 100 unique emails per request."),
        ("Clean, responsive UI: ", "Built with vanilla JavaScript, modern CSS, dynamic status badges, summary count cards, search filtering, and live counter."),
        ("CSV download: ", "Includes a download button to export the checked results and reasons into a CSV file.")
    ]
    for bold_prefix, text in features:
        p_feat = doc.add_paragraph(style='List Bullet')
        r_p = p_feat.add_run(bold_prefix)
        r_p.bold = True
        p_feat.add_run(text)
        p_feat.paragraph_format.space_after = Pt(4)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Deliverability Statuses")
    r_h3.font.color.rgb = primary_color
    r_h3.font.bold = True

    class_table = doc.add_table(rows=4, cols=3)
    class_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    class_widths = [Inches(1.8), Inches(2.3), Inches(2.4)]
    for row in class_table.rows:
        for idx, width in enumerate(class_widths):
            row.cells[idx].width = width

    headers = ["Status", "Technical Check", "Meaning"]
    for idx, name in enumerate(headers):
        cell = class_table.rows[0].cells[idx]
        set_cell_background(cell, "4F46E5")
        p = cell.paragraphs[0]
        r = p.add_run(name)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("Likely deliverable", "Valid syntax, working MX records, and mail server returned 250 OK on recipient check.", "The email is likely active, though spam filters or mailbox quotas can still cause issues."),
        ("Likely to bounce", "Malformed syntax, nonexistent domain (NXDOMAIN), missing MX/A records, or 550 rejection.", "Address is broken or rejected by the server. Sending to this will almost certainly bounce."),
        ("Unknown", "Domain has mail servers, but server timed out, greylisted the probe, or port 25 is blocked.", "Mailbox status cannot be confirmed without sending a message.")
    ]

    for row_idx, data in enumerate(rows_data):
        row = class_table.rows[row_idx + 1]
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Limitations and Real-World Considerations")
    r_h4.font.color.rgb = primary_color
    r_h4.font.bold = True

    p_lim = doc.add_paragraph(
        "1. Catch-All domains: Many corporate servers accept all incoming addresses initially (status 250 OK) and "
        "only drop invalid mailboxes afterwards, making it impossible to guarantee delivery from the handshake alone.\n"
        "2. Concealed status: Large providers like Microsoft, Yahoo, and iCloud hide recipient status to prevent spammers from harvesting active addresses.\n"
        "3. Free cloud hosting limitations: Free hosting providers like Render block outbound port 25. When hosted on Render, the app verifies syntax and DNS MX records, and marks valid domains as 'Unknown' with an explanation rather than throwing false errors."
    )
    p_lim.paragraph_format.space_after = Pt(12)

    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. Automated Testing")
    r_h5.font.color.rgb = primary_color
    r_h5.font.bold = True

    p_test = doc.add_paragraph(
        "I wrote a comprehensive test suite using pytest containing 22 tests covering:\n"
        "• Valid and invalid syntax formatting\n"
        "• Batch deduplication and the 100 email limit\n"
        "• Mocked DNS MX and fallback A record lookups, NXDOMAIN, and timeouts\n"
        "• Mocked SMTP response codes (250 deliverable, 550 bounce, 451 greylisting, timeouts)\n"
        "• Verification that the SMTP DATA command is never called\n"
        "• Flask routes (/api/check, /api/export-csv, /health)\n"
        "• CSV report format and content\n\n"
        "All 22 tests pass locally."
    )
    p_test.paragraph_format.space_after = Pt(12)

    doc.save(output_path)
    print(f"Generated docx at: {output_path}")


if __name__ == "__main__":
    create_submission_document()
